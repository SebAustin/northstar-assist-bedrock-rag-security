"""Fill the official Udacity ML-BOM and STRIDE-ML templates from YAML content files.

The templates in ``starter/`` are only ever read.  Content comes from ``content/mlbom.yaml`` and
``content/stride.yaml``; strings may carry ``{{TOKEN}}`` placeholders (see ``KNOWN_TOKENS``) that are
resolved from ``infra/state.env`` (``KEY=VALUE`` lines).  ``{{DATE}}`` is built in.  Unresolved tokens
are left in place and reported; ``--check`` turns them into a failure.

Usage::

    .venv/bin/python tools/fill_templates.py [--check]

The low-level helpers (``fill_kv_table``, ``fill_grid_table``, ``replace_placeholder_paragraph``,
``insert_paragraph_after``, ``insert_table_after``) work on a loaded ``docx.Document`` and keep the
template's own styles.  Inline ``**bold**`` markup is supported in every text value.
"""
from __future__ import annotations

import argparse
import copy
import re
import shlex
import sys
from datetime import date
from pathlib import Path
from typing import Any

import yaml
from content_schema import (  # noqa: F401  (re-exported)
    ASSET_HEADER,
    COMPONENT_HEADER,
    ID_PREFIX,
    MODEL_LABELS,
    STATUSES,
    STRIDE_CATEGORIES,
    ContentError,
    load_yaml,
    priority_for,
    risk_score,
    validate_mlbom,
    validate_stride,
)
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.shared import Pt, Twips
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent

KNOWN_TOKENS = ("HARNESS_ARN", "HARNESS_ROLE_NAME", "GATEWAY_ARN", "GATEWAY_ROLE_NAME", "KB_ID",
                "KB_ROLE_NAME", "BUCKET", "MODEL_ID", "GUARDRAIL_ARN")
TOKEN_RE = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
PLACEHOLDER_RE = re.compile(r"\[[^\]]*\]")
STATE_KEY_RE = re.compile(r"^[A-Z_][A-Z0-9_]*$")

PRIORITY_FILL = {"Critical": "F4B6B0", "High": "F9D3A8", "Medium": "FBEFA6", "Low": "CFE8CF"}
HEADER_FILL = "E7EBF0"
TABLE_WIDTH_DXA = 9000
DASH = "–"

THREAT_HEADER = ["Trust boundary", "Likelihood (1-5)", "Impact (1-5)", "Risk score", "Priority"]
REGISTER_HEADER = ["ID", "Threat", "STRIDE", "L", "I", "Score", "Priority", "Mitigation", "Status"]



# --------------------------------------------------------------------------- state and tokens


def load_state(path: Path) -> dict[str, str]:
    """Parse ``KEY=VALUE`` lines (shell-quoted values allowed); a missing file yields ``{}``."""
    path = Path(path)
    if not path.exists():
        return {}
    state: dict[str, str] = {}
    for line in path.read_text().splitlines():
        key, sep, raw = line.strip().partition("=")
        if not sep or key.startswith("#") or not STATE_KEY_RE.match(key):
            continue
        try:
            state[key] = " ".join(shlex.split(raw))
        except ValueError:
            print(f"warning: skipping unparsable state line for {key}", file=sys.stderr)
    return state


def resolve_tokens(data: Any, state: dict[str, str]) -> tuple[Any, set[str]]:
    """Return a copy of ``data`` with ``{{TOKEN}}`` replaced from ``state`` plus the unresolved names."""
    unresolved: set[str] = set()

    def walk(node: Any) -> Any:
        if isinstance(node, str):
            def sub(match: re.Match) -> str:
                name = match.group(1)
                if name in state:
                    return state[name]
                unresolved.add(name)
                return match.group(0)
            return TOKEN_RE.sub(sub, node)
        if isinstance(node, dict):
            return {key: walk(value) for key, value in node.items()}
        if isinstance(node, list):
            return [walk(item) for item in node]
        return node

    return walk(data), unresolved



# --------------------------------------------------------------------------- low-level docx helpers


def _is_heading(element: Any) -> bool:
    if element.tag != qn("w:p"):
        return False
    return any(v.startswith("Heading") for v in element.xpath("./w:pPr/w:pStyle/@w:val"))


def _wrap(element: Any, doc: Any) -> Paragraph:
    return Paragraph(element, doc._body)


def _clear_paragraph(p_element: Any) -> None:
    for child in list(p_element):
        if child.tag != qn("w:pPr"):
            p_element.remove(child)


def add_marked_runs(paragraph: Paragraph, text: str, *, bold: bool = False, size_pt: float | None = None) -> None:
    """Append runs for ``text``; segments between ``**`` markers are bold."""
    for index, chunk in enumerate(str(text).split("**")):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        run.bold = True if (bold or index % 2 == 1) else None
        if size_pt:
            run.font.size = Pt(size_pt)


def set_paragraph_text(paragraph: Paragraph, text: str, *, bold: bool = False) -> None:
    """Replace the runs of ``paragraph`` with ``text`` (the paragraph style is kept)."""
    _clear_paragraph(paragraph._p)
    add_marked_runs(paragraph, text, bold=bold)


def set_cell_text(cell: Any, text: Any, *, bold: bool = False, size_pt: float | None = None) -> None:
    """Replace the content of a table cell; ``\\n`` starts a new paragraph in the same style."""
    lines = ("" if text is None else str(text)).split("\n")
    first, *extra = cell.paragraphs
    for paragraph in extra:
        paragraph._p.getparent().remove(paragraph._p)
    _clear_paragraph(first._p)
    add_marked_runs(first, lines[0], bold=bold, size_pt=size_pt)
    previous = first
    for line in lines[1:]:
        clone = copy.deepcopy(first._p)
        _clear_paragraph(clone)
        previous._p.addnext(clone)
        previous = Paragraph(clone, cell)
        add_marked_runs(previous, line, bold=bold, size_pt=size_pt)


def _shade_cell(cell: Any, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _append_row_like(table: Table, template_row: _Row) -> _Row:
    clone = copy.deepcopy(template_row._tr)
    table.rows[-1]._tr.addnext(clone)
    row = _Row(clone, table)
    for cell in row.cells:
        set_cell_text(cell, "")
    return row


def fill_kv_table(table: Table, mapping: dict[str, Any], *, add_missing: bool = False) -> None:
    """Set the Value cell of every row whose Field cell matches a key of ``mapping``.

    Unknown fields raise ``ContentError`` unless ``add_missing`` is set, in which case a row is appended.
    """
    remaining = dict(mapping)
    for row in table.rows[1:]:
        key = row.cells[0].text.strip()
        if key in remaining:
            set_cell_text(row.cells[1], remaining.pop(key))
    if remaining and not add_missing:
        raise ContentError(f"fields not found in table: {sorted(remaining)}")
    for key, value in remaining.items():
        row = _append_row_like(table, table.rows[-1])
        set_cell_text(row.cells[0], key, bold=True)
        set_cell_text(row.cells[1], value)


def fill_grid_table(table: Table, rows: list[list[Any]]) -> None:
    """Fill the rows under the header with ``rows``.

    Existing template rows are reused (their labels are overwritten), unused ones are removed and
    missing ones are appended by copying the last row, so cell styles carry over.
    """
    ncols = len(table.rows[0].cells)
    for values in rows:
        if len(values) > ncols:
            raise ContentError(f"row has {len(values)} cells but the table has {ncols} columns: {values}")
    body = list(table.rows[1:])
    while len(body) < len(rows):
        body.append(_append_row_like(table, body[-1] if body else table.rows[0]))
    for row, values in zip(body, rows, strict=False):
        padded = [*values, *[""] * (ncols - len(values))]
        for cell, value in zip(row.cells, padded, strict=True):
            set_cell_text(cell, value)
    for row in body[len(rows):]:
        table._tbl.remove(row._tr)


# --------------------------------------------------------------------------- paragraph helpers


def find_paragraph(doc: Any, text: str, *, exact: bool = False) -> Paragraph:
    """First body paragraph whose text starts with (or, with ``exact``, equals) ``text``."""
    for paragraph in doc.paragraphs:
        candidate = paragraph.text.strip()
        if candidate == text if exact else candidate.startswith(text):
            return paragraph
    raise ContentError(f"paragraph not found: {text!r}")


def find_heading(doc: Any, text: str) -> Paragraph:
    """The heading paragraph whose text equals ``text``."""
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading") and paragraph.text.strip() == text:
            return paragraph
    raise ContentError(f"heading not found: {text!r}")


def table_after_heading(doc: Any, heading_text: str) -> Table:
    """The first table that follows the given heading and precedes the next one."""
    element = find_heading(doc, heading_text)._p.getnext()
    while element is not None and not _is_heading(element):
        if element.tag == qn("w:tbl"):
            return Table(element, doc._body)
        element = element.getnext()
    raise ContentError(f"no table under heading {heading_text!r}")


def find_table(doc: Any, header: list[str]) -> Table:
    """The first table whose header row equals ``header``."""
    for table in doc.tables:
        if [c.text.strip() for c in table.rows[0].cells] == header:
            return table
    raise ContentError(f"table not found with header {header}")


def _bullet_num_id(element: Any) -> int | None:
    found = element.getroottree().getroot().xpath(".//w:numPr/w:numId/@w:val")
    return int(found[0]) if found else None


def insert_paragraph_after(anchor: Paragraph | Table, text: str = "", style: str | None = None, *,
                           bold_prefix: str | None = None, bold: bool = False, bullet: bool = False,
                           keep_with_next: bool = False, size_pt: float | None = None) -> Paragraph:
    """Insert a paragraph directly after a paragraph or table and return it."""
    p_element = OxmlElement("w:p")
    anchor._element.addnext(p_element)
    paragraph = Paragraph(p_element, anchor._parent)
    if style:
        paragraph.style = style
    num_id = _bullet_num_id(p_element) if bullet else None
    if num_id is not None:
        num_pr = p_element.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num_id
    elif bullet:
        text = f"• {text}"
    if bold_prefix:
        add_marked_runs(paragraph, bold_prefix, bold=True, size_pt=size_pt)
        text = f" {text}"
    add_marked_runs(paragraph, text, bold=bold, size_pt=size_pt)
    if keep_with_next:
        paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _insert_many(anchor: Paragraph | Table, texts: list[str], style: str | None = None, **kwargs: Any) -> Any:
    """Insert one paragraph per text in order; returns the last element (or ``anchor`` if empty)."""
    for text in texts:
        anchor = insert_paragraph_after(anchor, text, style, **kwargs)
    return anchor


def replace_placeholder_paragraph(doc: Any, placeholder_prefix: str, new_texts: list[str],
                                  style: str | None = None) -> list[Paragraph]:
    """Replace the paragraph starting with ``placeholder_prefix`` by one paragraph per text, same style."""
    old = find_paragraph(doc, placeholder_prefix)
    style = style or old.style.name
    created: list[Paragraph] = []
    anchor: Paragraph = old
    for text in new_texts:
        anchor = insert_paragraph_after(anchor, text, style)
        created.append(anchor)
    old._p.getparent().remove(old._p)
    return created


def fill_labeled_paragraph(doc: Any, label: str, text: str | None = None, bullets: list[str] | None = None
                           ) -> Paragraph:
    """Keep the bold label run of ``Label: [placeholder]``; append ``text`` and optional bullet paragraphs."""
    paragraph = find_paragraph(doc, label)
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)
    if text:
        add_marked_runs(paragraph, f" {text}")
    _insert_many(paragraph, bullets or [], "Compact", bullet=True)
    return paragraph


def insert_table_after(doc: Any, anchor: Paragraph | Table, header: list[str], rows: list[list[Any]], *,
                       widths: list[float] | None = None, size_pt: float | None = None,
                       style: str = "Table") -> Table:
    """Insert a bordered table (template ``Table`` style, shaded bold header, repeating) after ``anchor``."""
    ncols = len(header)
    for row in rows:
        if len(row) != ncols:
            raise ContentError(f"row has {len(row)} cells, header has {ncols}: {row}")
    weights = widths or [1.0] * ncols
    if len(weights) != ncols:
        raise ContentError("widths must have one entry per column")
    dxa = [int(w * TABLE_WIDTH_DXA / sum(weights)) for w in weights]
    tbl = CT_Tbl.new_tbl(1 + len(rows), ncols, Twips(TABLE_WIDTH_DXA))
    anchor._element.addnext(tbl)
    table = Table(tbl, anchor._parent)
    table.style = style
    table.autofit = False
    _finish_table_properties(tbl)
    for grid_col, width in zip(tbl.tblGrid.findall(qn("w:gridCol")), dxa, strict=True):
        grid_col.set(qn("w:w"), str(width))
    for r_index, values in enumerate([header, *rows]):
        for c_index, value in enumerate(values):
            cell = table.cell(r_index, c_index)
            cell.width = Twips(dxa[c_index])
            cell.paragraphs[0].style = "Compact"
            set_cell_text(cell, value, bold=r_index == 0, size_pt=size_pt)
            if r_index == 0:
                _shade_cell(cell, HEADER_FILL)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


def _finish_table_properties(tbl: Any) -> None:
    """Full-width (pct) table with light horizontal rules, inserted in schema order."""
    tbl_pr = tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        side = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "BFBFBF")):
            side.set(qn(f"w:{key}"), value)
        borders.append(side)
    successor = next((tbl_pr.find(qn(tag)) for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
                      if tbl_pr.find(qn(tag)) is not None), None)
    if successor is not None:
        successor.addprevious(borders)
    else:
        tbl_pr.append(borders)


def _add_rule_after(doc: Any, anchor: Paragraph | Table) -> Paragraph:
    """Copy the template's horizontal-rule paragraph after ``anchor``."""
    rule = next((p for p in doc.element.body.iter(qn("w:p")) if p.find(f".//{qn('w:pict')}") is not None), None)
    clone = copy.deepcopy(rule) if rule is not None else OxmlElement("w:p")
    anchor._element.addnext(clone)
    return Paragraph(clone, anchor._parent)


def _section_last_element(heading: Paragraph) -> Any:
    """Last paragraph/table element of the section that starts at ``heading`` (bookmarks skipped)."""
    content_tags = (qn("w:p"), qn("w:tbl"))
    last = heading._p
    element = heading._p.getnext()
    while element is not None and not _is_heading(element):
        if element.tag in content_tags:
            last = element
        element = element.getnext()
    return last


def _document_texts(doc: Any) -> list[str]:
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


def find_placeholders(doc: Any) -> list[str]:
    """Texts (paragraphs and table cells) that still contain ``[...]`` placeholder text."""
    return [text for text in _document_texts(doc) if PLACEHOLDER_RE.search(text)]


def find_unresolved_tokens(doc: Any) -> set[str]:
    """Names of ``{{TOKEN}}`` placeholders still present in the document text."""
    return {name for text in _document_texts(doc) for name in TOKEN_RE.findall(text)}



# --------------------------------------------------------------------------- ML-BOM builder

MODEL_SECTION_TABLES = (
    ("1. Model Information", "kv", "model_info"),
    ("1.1 BOM Generation (Optional)", "kv", "bom_generation"),
    ("1.2 Other References (Optional)", "grid", "references"),
    ("2.1 Training Datasets", "grid", "datasets"),
    ("2.2 Architecture Details", "kv", "architecture"),
    ("2.3 Model Lineage", "grid", "lineage"),
    ("2.4 Input/Output", "kv", "io"),
    ("2.5 Hardware & Software (Optional)", "grid", "hardware_software"),
    ("2.6 Software Required for Execution", "kv", "software_required"),
)
USAGE_PLACEHOLDERS = (
    ("[Required: Description of how", "intended"),
    ("[Required: Description of use cases", "out_of_scope"),
    ("[Required: Description of what constitutes", "misuse"),
)


def _fill_model(doc: Any, model: dict[str, Any]) -> None:
    """Fill every table and placeholder of one (still unfilled) copy of template sections 1-3."""
    for heading, kind, key in MODEL_SECTION_TABLES:
        table = table_after_heading(doc, heading)
        if kind == "kv":
            fill_kv_table(table, model[key], add_missing=True)
        else:
            fill_grid_table(table, model[key])
    replace_placeholder_paragraph(doc, "[Tags or labels", [model["tags"]])
    for prefix, key in USAGE_PLACEHOLDERS:
        replace_placeholder_paragraph(doc, prefix, model["usage"][key])


def _model_block(doc: Any) -> list[Any]:
    """Elements of template sections 1-3 (bookmarks excluded), from '1. Model Information' to Attestation."""
    skip = (qn("w:bookmarkStart"), qn("w:bookmarkEnd"))
    block, element = [], find_heading(doc, "1. Model Information")._p
    while element is not None and not (_is_heading(element) and "".join(
            t.text or "" for t in element.iter(qn("w:t"))) == "Attestation"):
        if element.tag not in skip:
            block.append(element)
        element = element.getnext()
    return block


def _prefix_headings(doc: Any, block: list[Any], label: str) -> None:
    for element in block:
        if _is_heading(element):
            paragraph = _wrap(element, doc)
            set_paragraph_text(paragraph, f"Model {label} {DASH} {paragraph.text}")


def _add_banner(doc: Any, before: Any, label: str, model: dict[str, Any]) -> None:
    """Insert 'Model X: title' heading and summary directly before the element ``before``."""
    anchor = Paragraph(OxmlElement("w:p"), doc._body)
    before.addprevious(anchor._p)
    anchor.style = "Heading 2"
    set_paragraph_text(anchor, f"Model {label}: {model['title']}")
    insert_paragraph_after(anchor, model["summary"], "First Paragraph")


def _add_system_overview(doc: Any, overview: dict[str, Any]) -> None:
    rule = _wrap(_section_last_element(find_heading(doc, "Overview")), doc)
    last = insert_paragraph_after(rule, "System Overview", "Heading 2")
    last = _insert_many(last, overview["intro"], "Body Text")
    last = insert_paragraph_after(last, "Component data flow", "Heading 3")
    last = _insert_many(last, overview["dataflow"], "Body Text")
    for line in overview["flow"]:
        last = insert_paragraph_after(last, line, "Source Code", size_pt=9)
        last.paragraph_format.space_after = Pt(0)
        last.paragraph_format.space_before = Pt(0)
        for run in last.runs:
            run.font.name = "Consolas"
    last = insert_paragraph_after(last, "Request path", "Heading 3")
    last = _insert_many(last, [f"**{n}.** {text}" for n, text in enumerate(overview["request_path"], 1)],
                        "Compact")
    last = insert_paragraph_after(last, "Ingestion path", "Heading 3")
    last = _insert_many(last, overview["ingestion_path"], "Body Text")
    _add_rule_after(doc, last)


def _add_components_and_findings(doc: Any, after: Any, content: dict[str, Any]) -> None:
    components = content["supporting_components"]
    last = insert_paragraph_after(after, "4. Supporting Components and IAM Roles", "Heading 2")
    last = _insert_many(last, [components["intro"]], "Body Text")
    last = insert_table_after(doc, last, COMPONENT_HEADER, components["rows"],
                              widths=[1.3, 0.9, 0.9, 1.5, 1.7, 1.7], size_pt=8)
    last = _add_rule_after(doc, last)
    last = insert_paragraph_after(last, "5. Findings", "Heading 2")
    last = insert_paragraph_after(last, content["findings_intro"], "Body Text")
    for finding in content["findings"]:
        title = f"**{finding['id']}: {finding['title']}** ({finding['severity']})"
        last = insert_paragraph_after(last, title, "Body Text", keep_with_next=True)
        last = insert_paragraph_after(last, finding["text"], "Compact")
        last = insert_paragraph_after(last, finding["action"], "Compact", bold_prefix="Action:")
    _add_rule_after(doc, last)


def render_mlbom(content: dict[str, Any], template: Path, state: dict[str, str] | None = None
                 ) -> tuple[Any, set[str]]:
    """Build the ML-BOM in memory; returns the document and the ``{{TOKEN}}`` names left unresolved."""
    validate_mlbom(content)
    resolved, _ = resolve_tokens(content, {"DATE": date.today().isoformat(), **(state or {})})
    doc = Document(str(template))
    block = _model_block(doc)
    pristine = [copy.deepcopy(element) for element in block]
    model_a, model_b = resolved["models"]
    _fill_model(doc, model_a)
    _prefix_headings(doc, block, MODEL_LABELS[0])
    _add_banner(doc, block[0], MODEL_LABELS[0], model_a)
    for element in reversed(pristine):
        block[-1].addnext(element)
    _fill_model(doc, model_b)
    _prefix_headings(doc, pristine, MODEL_LABELS[1])
    _add_banner(doc, pristine[0], MODEL_LABELS[1], model_b)
    _add_system_overview(doc, resolved["system_overview"])
    _add_components_and_findings(doc, _wrap(pristine[-1], doc), resolved)
    fill_grid_table(table_after_heading(doc, "Attestation"), resolved["attestation"])
    _assert_no_placeholders(doc)
    return doc, find_unresolved_tokens(doc)


# --------------------------------------------------------------------------- STRIDE-ML builder


def _fill_system_overview(doc: Any, system: dict[str, Any]) -> None:
    fill_labeled_paragraph(doc, "System Name:", system["name"])
    fill_labeled_paragraph(doc, "Purpose:", system["purpose"])
    fill_labeled_paragraph(doc, "Architecture:", system["architecture"])
    fill_labeled_paragraph(doc, "Key Components:", bullets=system["key_components"])
    boundaries = [f"**{b['id']} {DASH} {b['name']}:** {b['text']}" for b in system["trust_boundaries"]]
    fill_labeled_paragraph(doc, "Trust Boundaries:", bullets=boundaries)


def _score_rows(threats: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [{**t, "score": risk_score(t["likelihood"], t["impact"]),
             "priority": priority_for(risk_score(t["likelihood"], t["impact"]))} for t in threats]


def _color_priority(table: Table, column: int, first_row: int = 1) -> None:
    for row in table.rows[first_row:]:
        cell = row.cells[column]
        _shade_cell(cell, PRIORITY_FILL[cell.text.strip()])


def _add_assets_and_register(doc: Any, content: dict[str, Any], threats: list[dict[str, Any]]) -> None:
    assets = content["assets"]
    paragraphs = replace_placeholder_paragraph(doc, "[Describe the sensitive data", [assets["intro"]])
    table = insert_table_after(doc, paragraphs[-1], ASSET_HEADER, assets["rows"],
                               widths=[1.2, 1.0, 1.6, 2.2, 0.8], size_pt=8.5)
    rule_element = table._tbl.getnext()
    heading = insert_paragraph_after(_wrap(rule_element, doc), "Prioritized Risk Register", "Heading 2")
    intro = insert_paragraph_after(heading, content["register_intro"], "Body Text")
    ordered = sorted(threats, key=lambda t: (-t["score"], t["id"]))
    rows = [[t["id"], t["name"], STRIDE_CATEGORIES[t["category"]], t["likelihood"], t["impact"], t["score"],
             t["priority"], t["mitigation_summary"], t["status"]] for t in ordered]
    register = insert_table_after(doc, intro, REGISTER_HEADER, rows,
                                  widths=[0.55, 1.5, 0.95, 0.3, 0.3, 0.6, 0.7, 2.05, 0.85], size_pt=7.5)
    _color_priority(register, 6)
    _add_rule_after(doc, register)


def _add_threat(doc: Any, anchor: Any, threat: dict[str, Any], boundary_names: dict[str, str]) -> Any:
    last = insert_paragraph_after(anchor, f"{threat['id']}: {threat['name']}", "Body Text", bold=True,
                                  keep_with_next=True)
    last = insert_paragraph_after(last, threat["scenario"], "Body Text", bold_prefix="Scenario:",
                                  keep_with_next=True)
    boundary = f"{threat['boundary']}: {boundary_names[threat['boundary']]}"
    row = [boundary, threat["likelihood"], threat["impact"], threat["score"], threat["priority"]]
    table = insert_table_after(doc, last, THREAT_HEADER, [row], widths=[2.4, 1.1, 1.0, 1.0, 1.0], size_pt=9)
    _color_priority(table, 4)
    last = insert_paragraph_after(table, threat["mitigation"], "Body Text", bold_prefix="Mitigation:")
    return insert_paragraph_after(last, threat["residual"], "Body Text", bold_prefix="Residual:")


def _fill_threat_sections(doc: Any, content: dict[str, Any], threats: list[dict[str, Any]]) -> None:
    names = {b["id"]: b["name"] for b in content["system"]["trust_boundaries"]}
    for number, (key, label) in enumerate(STRIDE_CATEGORIES.items(), start=1):
        heading = find_heading(doc, f"3.{number} {label}")
        intro, title, description, mitigation = _threat_template_paragraphs(doc, heading)
        for stale in (title, description, mitigation):
            stale._p.getparent().remove(stale._p)
        set_paragraph_text(intro, content["category_intros"][key])
        anchor: Any = intro
        for threat in (t for t in threats if t["category"] == key):
            anchor = _add_threat(doc, anchor, threat, names)


def _threat_template_paragraphs(doc: Any, heading: Paragraph) -> list[Paragraph]:
    paragraphs, element = [], heading._p.getnext()
    while len(paragraphs) < 4 and element is not None:
        if element.tag == qn("w:p"):
            paragraphs.append(_wrap(element, doc))
        element = element.getnext()
    texts = [p.text for p in paragraphs]
    if len(texts) < 4 or not texts[0].startswith("[") or not texts[1].startswith("[Threat ID]") \
            or not texts[3].startswith("Mitigation:"):
        raise ContentError(f"unexpected threat block under {heading.text!r}: {texts}")
    return paragraphs


def render_stride(content: dict[str, Any], template: Path, state: dict[str, str] | None = None
                  ) -> tuple[Any, set[str]]:
    """Build the STRIDE-ML document in memory; returns it and the ``{{TOKEN}}`` names left unresolved."""
    validate_stride(content)
    resolved, _ = resolve_tokens(content, {"DATE": date.today().isoformat(), **(state or {})})
    threats = _score_rows(resolved["threats"])
    doc = Document(str(template))
    if resolved.get("title"):
        set_paragraph_text(find_paragraph(doc, "STRIDE-ML Threat Model Template", exact=True), resolved["title"])
    _fill_system_overview(doc, resolved["system"])
    _add_assets_and_register(doc, resolved, threats)
    _fill_threat_sections(doc, resolved, threats)
    replace_placeholder_paragraph(doc, "[Describe risks that remain", [resolved["residual_intro"]])
    replace_placeholder_paragraph(doc, "[Risk Name]",
                                  [f"**{r['name']}:** {r['text']}" for r in resolved["residual_risks"]])
    recommendations = [f"**R{n}.** {text}" for n, text in enumerate(resolved["recommendations"], 1)]
    replace_placeholder_paragraph(doc, "[Provide actionable",
                                  [resolved["recommendations_intro"], *recommendations])
    fill_grid_table(find_table(doc, ["Role", "Name", "Date"]), resolved["signoff"])
    _assert_no_placeholders(doc)
    return doc, find_unresolved_tokens(doc)


# --------------------------------------------------------------------------- build entry points and CLI


def _assert_no_placeholders(doc: Any) -> None:
    leftovers = find_placeholders(doc)
    if leftovers:
        raise ContentError(f"placeholder text remains in the output: {leftovers[:3]}")


def _save(doc: Any, template: Path, out: Path) -> None:
    out = Path(out)
    if out.resolve() == Path(template).resolve():
        raise ContentError("refusing to overwrite the starter template")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def build_mlbom(content: dict[str, Any], template: Path, out: Path, state: dict[str, str] | None = None
                ) -> set[str]:
    """Render and save the ML-BOM; returns unresolved token names."""
    doc, unresolved = render_mlbom(content, template, state)
    _save(doc, template, out)
    return unresolved


def build_stride(content: dict[str, Any], template: Path, out: Path, state: dict[str, str] | None = None
                 ) -> set[str]:
    """Render and save the STRIDE-ML document; returns unresolved token names."""
    doc, unresolved = render_stride(content, template, state)
    _save(doc, template, out)
    return unresolved


def _parse_args(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fill the ML-BOM and STRIDE-ML templates from YAML content.")
    parser.add_argument("--check", action="store_true", help="exit 1 if any {{TOKEN}} remains unresolved")
    parser.add_argument("--state", type=Path, default=ROOT / "infra" / "state.env")
    parser.add_argument("--content-dir", type=Path, default=ROOT / "content")
    parser.add_argument("--starter-dir", type=Path, default=ROOT / "starter")
    parser.add_argument("--out-dir", type=Path, default=ROOT / "deliverables")
    args = parser.parse_args(argv)
    starter = args.starter_dir.resolve()
    if args.out_dir.resolve() == starter or starter in args.out_dir.resolve().parents:
        parser.error("--out-dir must not be the starter directory")
    return args


def main(argv: list[str] | None = None) -> int:
    """CLI entry point; returns the process exit code."""
    args = _parse_args(argv)
    state = load_state(args.state)
    jobs = (
        ("mlbom.yaml", "ML-BOM Template.docx", "ML-BOM.docx", render_mlbom),
        ("stride.yaml", "STRIDE-ML Template.docx", "STRIDE-ML.docx", render_stride),
    )
    unresolved: set[str] = set()
    rendered = []
    try:
        for content_name, template_name, out_name, render in jobs:
            template = args.starter_dir / template_name
            doc, missing = render(load_yaml(args.content_dir / content_name), template, state)
            unresolved |= missing
            rendered.append((doc, template, args.out_dir / out_name))
    except (ContentError, OSError, yaml.YAMLError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    if unresolved:
        names = ", ".join(sorted(unresolved))
        print(f"warning: unresolved tokens left in the output: {names}", file=sys.stderr)
        if args.check:
            print("error: --check failed, nothing written", file=sys.stderr)
            return 1
    for doc, template, out in rendered:
        _save(doc, template, out)
        print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
